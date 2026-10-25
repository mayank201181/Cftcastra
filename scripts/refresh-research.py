"""Public research feeds. Preserve old observations on failure, and never reset their dates.
Dependencies: openpyxl for JPX, pdftotext (poppler-utils) for official CME PDFs.
No credentials, paid subscriptions, portfolio records or order execution.
"""
import concurrent.futures, csv, datetime as dt, gzip, hashlib, io, json, math, pathlib, re, subprocess, tempfile, urllib.parse, urllib.request
from zoneinfo import ZoneInfo
import openpyxl

ROOT=pathlib.Path(__file__).resolve().parents[1]
DATA=ROOT/'data'
NOW=dt.datetime.now(dt.timezone.utc)
TODAY=NOW.date()
CONFIG=json.loads((DATA/'research-markets.json').read_text())

def fetch(url):
    req=urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0 (CFTC Astra public research)','Accept':'*/*'})
    with urllib.request.urlopen(req,timeout=45) as r:return r.read()

def read_json(path,default):
    if not path.exists():return default
    b=path.read_bytes()
    return json.loads(gzip.decompress(b) if path.suffix=='.gz' else b)

def write(path,value):
    b=json.dumps(value,separators=(',',':'),allow_nan=False).encode()
    if path.suffix=='.gz':b=gzip.compress(b,mtime=0)
    tmp=path.with_suffix('.tmp');tmp.write_bytes(b);tmp.replace(path)

def num(value):
    try:
        x=float(str(value).strip().replace(',',''))
        return x if math.isfinite(x) else None
    except (TypeError,ValueError):return None

def datestr(value):return dt.datetime.strptime(value.strip(),'%y%m%d').date().isoformat()

def prices():
    path=DATA/'research-prices.json.gz';old=read_json(path,{'markets':{}});out=dict(old['markets']);issues=[]
    jobs=list(CONFIG)+[{**m,'code':'bt:'+m['code'],'symbol':m['backtest'],'kind':'ETF adjusted close','quote':'USD per ETF share','inverse':False} for m in CONFIG if m.get('backtest')]
    def one(m):
        prior=out.get(m['code'],{});history=prior.get('history',[])
        # Corporate actions can rescale the entire adjusted-close series. Never splice adjustment bases.
        start=(dt.date.fromisoformat(history[-1][0])-dt.timedelta(days=14)) if history and m['kind']!='ETF adjusted close' else dt.date(2006 if m['kind']!='nearby futures' else 2021,1,1)
        end=TODAY+dt.timedelta(days=1)
        args={'period1':int(dt.datetime.combine(start,dt.time(),dt.timezone.utc).timestamp()),'period2':int(dt.datetime.combine(end,dt.time(),dt.timezone.utc).timestamp()),'interval':'1d'}
        url='https://query2.finance.yahoo.com/v8/finance/chart/'+urllib.parse.quote(m['symbol'],safe='')+'?'+urllib.parse.urlencode(args)
        j=json.loads(fetch(url))['chart']['result'][0];tz=ZoneInfo(j['meta'].get('exchangeTimezoneName','UTC'))
        vals=j['indicators'].get('adjclose',[{}])[0].get('adjclose') if m['kind']=='ETF adjusted close' else j['indicators']['quote'][0]['close']
        if vals is None:raise ValueError('Adjusted close unavailable')
        # Daily bars stamped at the session open must not be mistaken for completed closes.
        cutoff=TODAY if NOW.hour>=23 else TODAY-dt.timedelta(days=1)
        is_fx=m['sector']=='Foreign exchange'
        merged={} if m['kind']=='ETF adjusted close' else {p[0]:p[1] for p in history if not is_fx or dt.date.fromisoformat(p[0]).weekday()<5}
        for stamp,v in zip(j.get('timestamp',[]),vals):
            date=dt.datetime.fromtimestamp(stamp,tz).date()
            if v is None or not math.isfinite(v) or v<=0 or date>cutoff or (is_fx and date.weekday()>=5):continue
            merged[date.isoformat()]=round(1/v if m.get('inverse') else v,10)
        h=sorted(merged.items())
        if not h or len(h)<121:raise ValueError('Not enough complete daily observations')
        return {**m,'history':h,'asof':h[-1][0],'source':'Yahoo Finance daily close proxy','sourceUrl':'https://finance.yahoo.com/quote/'+urllib.parse.quote(m['symbol'],safe='')+'/history/','retrievedAt':NOW.isoformat()}
    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
        futures={pool.submit(one,m):m for m in jobs}
        for f,m in futures.items():
            try:r=f.result();out[m['code']]=r;print('PRICE',m['code'],len(r['history']),r['asof'],flush=True)
            except Exception as e:issues.append({'source':'price:'+m['code'],'error':str(e)[:180]});print('PRICE ERROR',m['code'],str(e),flush=True)
    write(path,{'retrievedAt':NOW.isoformat(),'markets':out,'issues':issues});return issues

def ice(previous):
    by={(r['code'],r['basis'],r['date']):r for r in previous}
    years=[TODAY.year] if previous else range(2016,TODAY.year+1)
    def one(year):
        url=f'https://www.ice.com/publicdocs/futures/COTHist{year}.csv';rows=[]
        for r in csv.DictReader(io.StringIO(fetch(url).decode('utf-8-sig'))):
            name=r['Market_and_Exchange_Names'];code='ice:B' if 'Brent Crude' in name else 'ice:G' if 'Gasoil' in name else None
            if not code:continue
            date=datestr(r['As_of_Date_In_Form_YYMMDD']);basis=r['FutOnly_or_Combined'].strip();oi=num(r['Open_Interest_All']);long=num(r['M_Money_Positions_Long_All']);short=num(r['M_Money_Positions_Short_All'])
            if date>TODAY.isoformat() or basis not in ('FutOnly','Combined') or oi is None or oi<=0 or long is None or short is None or min(long,short)<0:raise ValueError('Invalid ICE row')
            rows.append({'code':code,'name':'Brent' if code=='ice:B' else 'Gasoil','date':date,'basis':basis,'group':'Managed money','oi':oi,'long':long,'short':short,'net':long-short,'netoi':100*(long-short)/oi,'units':r['Contract_Units'],'sourceUrl':url})
        return rows
    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
        for result in pool.map(one,years):
            for r in result:by[(r['code'],r['basis'],r['date'])]=r
    if not by:raise ValueError('No ICE records')
    return sorted(by.values(),key=lambda r:(r['code'],r['basis'],r['date']))

def ttf(previous):
    url='https://registers.esma.europa.eu/solr/esma_registers_coder58/select?'+urllib.parse.urlencode({'q':'tv_mic:NDEX AND wp_vpcode:TFM','rows':2000,'wt':'json','sort':'wp_refdt asc,wp_pubdt asc'})
    result=json.loads(fetch(url))['response']
    if result['numFound']>2000:raise ValueError('ESMA pagination required; preserve existing archive')
    by={(r['date'],r['basis'],r['units']):r for r in previous}
    for r in result['docs']:
        if r.get('wp_sts') not in ('NEWT','AMND'):continue
        names=r.get('wp_pos_name',[]);cats=r.get('wp_pos_cat',[])
        indices=[i for i,(n,c) in enumerate(zip(names,cats)) if n=='Number of positions' and c=='TOTAL']
        if len(indices)!=1:continue
        i=indices[0]
        def field(k):return num(r.get(k,[])[i]) if len(r.get(k,[]))>i else None
        lo=field('wp_ifu_lo');sh=field('wp_ifu_sh')
        longs=[field('wp_'+g+'_lo') for g in ['ifi','ifu','oth','com','ope']];shorts=[field('wp_'+g+'_sh') for g in ['ifi','ifu','oth','com','ope']]
        if lo is None or sh is None or any(v is None or v<0 for v in longs+shorts):continue
        a,b=sum(longs),sum(shorts)
        if a<=0 or b<=0 or abs(a-b)/max(a,b)>.01:continue
        oi=(a+b)/2;date=r['wp_refdt'][:10];basis=r.get('wp_rptype','COMB');units=r.get('wp_munit')
        if date>TODAY.isoformat() or basis not in ('COMB','FUTR'):continue
        by[(date,basis,units)]={'code':'esma:TFM','name':'TTF · ICE Endex','date':date,'publishedAt':r['wp_pubdt'],'basis':basis,'units':units,'group':'Investment funds','long':lo,'short':sh,'net':lo-sh,'oi':oi,'netoi':100*(lo-sh)/oi,'oiNote':'Derived from mean of reported total long/short positions; rounding may differ','sourceUrl':url,'recordId':r['id']}
    if not by:raise ValueError('No valid ICE Endex TTF records')
    return sorted(by.values(),key=lambda r:(r['basis'],r['units'],r['date']))

def jpx(previous):
    page='https://www.jpx.co.jp/english/markets/statistics-equities/investor-type/index.html'
    html=fetch(page).decode();links=list(dict.fromkeys(re.findall(r'href="([^"]+stock_1_w_\d{8}_\d{8}\.xlsx)"',html)))
    if not links:raise ValueError('JPX spreadsheet links not found')
    by={r['date']:r for r in previous}
    for link in links[:5]:
        start,end=re.search(r'stock_1_w_(\d{8})_(\d{8})',link).groups();date=dt.datetime.strptime(end,'%Y%m%d').date().isoformat()
        url=urllib.parse.urljoin(page,link);book=openpyxl.load_workbook(io.BytesIO(fetch(url)),data_only=True);s=book.worksheets[0]
        if '1,000 yen' not in str(s['C7'].value):raise ValueError('JPX units changed')
        ri=next((r for r in range(1,s.max_row) if 'Tokyo & Nagoya' in str(s.cell(r,2).value)),None)
        if ri is None or 'Value' not in str(s.cell(ri+1,3).value):raise ValueError('JPX market schema changed')
        ri+=1
        groups=[('Foreign investors',[20,24]),('Domestic individuals',[12,16]),('Investment trusts',[32]),('Business corporations',[36]),('Trust banks',[52]),('Proprietary',[4,8])]
        values=[]
        for name,cols in groups:
            sales=buy=net=0
            for col in cols:
                a,b,c=[num(s.cell(ri,col+j).value) for j in range(3)]
                if None in (a,b,c) or abs((b-a)-c)>2:raise ValueError('JPX buy/sell/balance reconciliation failed')
                sales+=a;buy+=b;net+=c
            values.append({'group':name,'buyBn':buy/1e6,'sellBn':sales/1e6,'netBn':net/1e6})
        by[date]={'date':date,'start':dt.datetime.strptime(start,'%Y%m%d').date().isoformat(),'market':'Tokyo & Nagoya cash equities','units':'JPY billion','groups':values,'sourceUrl':url}
    return sorted(by.values(),key=lambda r:r['date'])

def pdf_text(url):
    b=fetch(url)
    if not b.startswith(b'%PDF'):raise ValueError('Expected PDF')
    with tempfile.TemporaryDirectory() as t:
        p=pathlib.Path(t)/'report.pdf';p.write_bytes(b)
        text=subprocess.check_output(['pdftotext','-layout',str(p),'-'],text=True)
    match=re.search(r'(?:Mon|Tue|Wed|Thu|Fri|Sat|Sun),\s+([A-Z][a-z]{2}\s+\d{1,2},\s+\d{4})',text)
    if not match:raise ValueError('PDF report date missing')
    date=dt.datetime.strptime(match[1],'%b %d, %Y').date().isoformat()
    if date>TODAY.isoformat():raise ValueError('Future PDF date')
    return text,date,hashlib.sha256(b).hexdigest()

def cme(previous):
    url='https://www.cmegroup.com/daily_bulletin/current/Section01B_Summary_Volume_And_Open_Interest_FX_Futures_And_Options.pdf'
    text,date,sha=pdf_text(url);page=text.split('\f')[0]
    header=next(l for l in page.splitlines() if 'OPEN INTEREST & CHGE' in l);anchor=header.index('OPEN INTEREST & CHGE')
    by={(r['code'],r['date']):r for r in previous};count=0
    for m in CONFIG:
        if not m.get('cme'):continue
        line=next((l for l in page.splitlines() if l.startswith(m['cme']+' ')),None)
        if not line:continue
        candidates=[x for x in re.finditer(r'\d+',line) if anchor-8<=x.start() and x.end()<=anchor+10]
        if len(candidates)!=1:raise ValueError('Ambiguous CME OI column '+m['cme'])
        token=candidates[0];oi=int(token.group());tail=line[token.end():header.rfind('VOLUME')];chg=re.search(r'([+-])\s*(\d+)',tail)
        change=int(chg[2])*(1 if chg[1]=='+' else -1) if chg else 0
        if oi<=0 or oi-change<=0:raise ValueError('Invalid CME OI')
        by[(m['code'],date)]={'code':m['code'],'name':m['name'],'date':date,'oi':oi,'change':change,'exchangeSymbol':m['cme'],'status':'Preliminary' if 'PRELIMINARY' in page else 'Final','sourceUrl':url,'sourceSha256':sha,'scope':'Standard contract, all listed maturities; excludes separate micro/mini contracts','noCot':m.get('noCot',False)};count+=1
    if count<8:raise ValueError('CME FX coverage unexpectedly low')
    return sorted(by.values(),key=lambda r:(r['code'],r['date']))

def parse_aud_strikes(text,date,url,sha):
    """Only publish ADU monthly buckets whose individual strike OI reconciles to CME TOTAL.
    Multiple series can share a strike/month in this bulletin; these are buckets, not exact expiries.
    """
    side=None;active=None;rows=[];out=[]
    for line in text.splitlines():
        if 'ADDITIONAL AUSSIE & NEW ZEALAND CALL' in line:side='call'
        if 'ADDITIONAL AUSSIE & NEW ZEALAND PUT' in line:side='put'
        header=re.match(r'^([A-Z]{3}\d{2})\s+(.+?)\s+\(FUTURES SETT',line)
        if header:
            active=header[1] if header[2].strip()=='ADU OPT' and side else None;rows=[];continue
        if not active:continue
        if re.match(r'^\d{4}\s',line):
            tokens=line.split();tail=tokens[:-2]
            unchanged=[i for i,t in enumerate(tokens) if t=='UNCH' and i>=len(tokens)-4]
            if unchanged:oi=tokens[unchanged[-1]-1]
            elif len(tail)>3 and tail[-2] in ('+','-'):oi=tail[-3]
            elif tail[-2]=='----':oi='----' # expired/exercised series, zero remaining OI
            else:raise ValueError('Unrecognized ADU option OI layout')
            if oi!='----' and not oi.isdigit():raise ValueError('Invalid ADU option OI')
            rows.append({'strike':int(tokens[0])/10000,'oi':0 if oi=='----' else int(oi)})
        elif re.match(r'^TOTAL\s+\d',line):
            tokens=line.split();total=int(tokens[2]);month=dt.datetime.strptime(active,'%b%y').strftime('%Y-%m')
            if sum(r['oi'] for r in rows)!=total:raise ValueError('ADU strike OI does not reconcile to monthly total')
            if month>date[:7]:
                by={}
                for r in rows:by[r['strike']]=by.get(r['strike'],0)+r['oi']
                out.append({'code':'232741','name':'AUD','date':date,'month':month,'side':side,'product':'ADU','totalOi':total,'strikes':[{'strike':k,'oi':v} for k,v in sorted(by.items()) if v>0],'status':'Preliminary' if 'PRELIMINARY' in text else 'Final','sourceUrl':url,'sourceSha256':sha,'scope':'ADU bulletin monthly bucket; may combine series with different expiries. Not all AUD options.'})
            active=None;rows=[]
    if len(out)<4 or not any(r['side']=='put' for r in out):raise ValueError('ADU option coverage unexpectedly low')
    return out

def strikes(previous):
    url='https://www.cmegroup.com/daily_bulletin/current/Section38_Australian_Dollar_New_Zealand_Dollar_Options.pdf'
    text,date,sha=pdf_text(url)
    return parse_aud_strikes(text,date,url,sha)

def refresh():
    price_issues=prices();path=DATA/'research-inputs.json';old=read_json(path,{});out=dict(old);statuses=[]
    jobs={'ice':ice,'ttf':ttf,'jpx':jpx,'cme':cme,'strikes':strikes}
    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
        tasks={name:pool.submit(fn,old.get(name,[])) for name,fn in jobs.items()}
        for name,f in tasks.items():
            try:
                result=f.result();out[name]=result;statuses.append({'source':name,'state':'ok','checkedAt':NOW.isoformat(),'asof':max(r['date'] for r in result),'records':len(result)});print('FEED',name,len(result),statuses[-1]['asof'],flush=True)
            except Exception as e:
                statuses.append({'source':name,'state':'error','checkedAt':NOW.isoformat(),'error':str(e)[:220],'retainedRecords':len(old.get(name,[]))});print('FEED ERROR',name,str(e),flush=True)
    statuses.append({'source':'reuters','state':'access-required','checkedAt':NOW.isoformat(),'error':'Current numerical poll is not available from an authorized public machine-readable feed. Import a dated survey CSV; no stale or invented scores are shown.','sourceUrl':'https://www.reuters.com/markets/currencies/'})
    out.update({'retrievedAt':NOW.isoformat(),'statuses':statuses,'priceIssues':price_issues});write(path,out)
    if not any(s['state']=='ok' for s in statuses):raise RuntimeError('All public research feeds failed; retained data preserved')

if __name__=='__main__':refresh()
